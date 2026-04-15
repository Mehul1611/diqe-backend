import asyncio
import csv
import logging
import os
import pytesseract
import fitz
from PIL import Image
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

logger = logging.getLogger(__name__)


class DataLoader:
    def __init__(self, input_data: dict):
        self.input_data = input_data
        self.model_id = input_data["model_id"]

    def _read_csv(self, path: str) -> str:
        content = []
        with open(path, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            for row in reader:
                content.append(", ".join(row))
        return "\n".join(content)

    def _read_txt(self, path: str) -> str:
        with open(path, encoding="utf-8") as f:
            return f.read()

    def _read_doc(self, path: str) -> str:
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)

    def _read_pdf(self, path: str) -> str:
        with open(path, "rb") as pdf:
            reader = PdfReader(pdf)
            text = "".join(page.extract_text() or "" for page in reader.pages)

        if len(text.strip()) < 20:
            ocr_text = self._ocr_pdf(path)
            if ocr_text.strip():
                return ocr_text
        return text

    def _ocr_image_path(self, path: str) -> str:
        if not pytesseract or not Image:
            logger.warning("OCR dependencies not installed; skipping OCR for %s", path)
            return ""
        try:
            img = Image.open(path)
            return pytesseract.image_to_string(img) or ""
        except Exception as exc:
            logger.warning("OCR failed for image %s: %s", path, exc)
            return ""

    def _ocr_pdf(self, path: str) -> str:
        if not fitz:
            logger.warning("PyMuPDF not installed; skipping PDF OCR for %s", path)
            return ""
        if not pytesseract or not Image:
            logger.warning("OCR dependencies not installed; skipping PDF OCR for %s", path)
            return ""

        out_parts: list[str] = []
        try:
            doc = fitz.open(path)
            for page in doc:
                pix = page.get_pixmap(dpi=200)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                out_parts.append(pytesseract.image_to_string(img) or "")
            doc.close()
        except Exception as exc:
            logger.warning("PDF OCR failed for %s: %s", path, exc)
            return ""
        return "\n".join(p for p in out_parts if p.strip())

    def _read_ppt(self, path: str) -> str:
        pres = Presentation(path)
        slides = []
        for slide in pres.slides:
            slide_text = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            slides.append("\n".join(slide_text))
        return "\n\n".join(slides)

    def _get_file_handler(self, extension: str):
        return {
            ".pdf": self._read_pdf,
            ".csv": self._read_csv,
            ".doc": self._read_doc,
            ".docx": self._read_doc,
            ".ppt": self._read_ppt,
            ".pptx": self._read_ppt,
            ".txt": self._read_txt,
            ".md": self._read_txt,
            ".png": self._ocr_image_path,
            ".jpg": self._ocr_image_path,
            ".jpeg": self._ocr_image_path,
            ".webp": self._ocr_image_path,
        }.get(extension)

    def _extract_file_sync(self, item: dict) -> dict | None:
        file_path = item.get("file_path", "")
        if not os.path.exists(file_path):
            logger.warning("File not found: %s", file_path)
            return None

        extension = os.path.splitext(file_path)[1].lower()
        handler = self._get_file_handler(extension)
        if not handler:
            logger.warning("Unsupported file type: %s", file_path)
            return None

        text = handler(file_path)
        return {
            "doc_id": item.get("doc_id", os.path.basename(file_path)),
            "title": item.get("title", os.path.basename(file_path)),
            "text": text,
        }

    async def _extract_file(self, item: dict) -> dict | None:
        try:
            return await asyncio.to_thread(self._extract_file_sync, item)
        except Exception as e:
            logger.error("Failed to process %s: %s", item.get("file_path"), e)
            return None

    async def extract_data(self):
        logger.info("Starting streaming text extraction for model_id=%s", self.model_id)
        for item in self.input_data.get("files_data", []):
            try:
                doc = await self._extract_file(item)
                if doc:
                    yield doc
            except Exception as e:
                logger.error("Error processing %s: %s", item.get("file_path"), e)
        logger.info("Extraction finished for model_id=%s", self.model_id)
